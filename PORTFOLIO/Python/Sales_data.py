import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import load_workbook
import seaborn as sns

my_path: str = r'C:\Users\mikko\OneDrive\Työpöytä\Sales_data.xlsx'
my_wb = load_workbook(my_path,data_only= True)
my_sheet = my_wb['Sheet1']
my_table = my_sheet.tables["Sales_data"]

table_cells = my_table.ref
#next we will create a list expression from table cells:
listed_cells = table_cells.split(':')
#Lets assess the columns and rows

start_col = ''.join(filter(str.isalpha, listed_cells[0]))
start_row = ''.join(filter(str.isnumeric, listed_cells[0]))
#then the end column and the row
end_col = ''.join(filter(str.isalpha, listed_cells[1]))
end_row = ''.join(filter(str.isnumeric, listed_cells[1]))

columns= f'{start_col}:{end_col}'
n_rows= int (end_row) - int (start_row)

My_range=pd.read_excel(my_path,
              sheet_name= "Sheet1",
              usecols= columns,
              nrows= n_rows,
              engine= "openpyxl"
              )
              
df=pd.DataFrame(My_range)
sns.set_style('whitegrid')
ax=sns.barplot(data= df, x = "City", y="Sales")
ax.tick_params(axis = 'x', labelsize=8, rotation=30)

avg=df['Sales'].mean()

ax.axhline(avg,color="green",linestyle='--')
ax.text(
    x=0.02,
    y=avg+(avg*0.01),
    s=f'{avg:.0f} Items',
    color='green',
    va='bottom',
    ha='center'
)

plt.show()
